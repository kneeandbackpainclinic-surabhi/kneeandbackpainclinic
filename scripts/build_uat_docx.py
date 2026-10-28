import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_uat_doc():
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color palette
    COLOR_PRIMARY = RGBColor(20, 48, 38)      # Forest Green
    COLOR_SECONDARY = RGBColor(194, 94, 59)   # Warm Terracotta
    COLOR_ACCENT = RGBColor(16, 185, 129)     # Emerald Green
    COLOR_TEXT = RGBColor(33, 43, 54)         # Charcoal Slate
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray
    COLOR_HEADER_BG = "E8F5E9"               # Soft Mint
    COLOR_ZEBRA_BG = "F8FAF9"                # Off White / Mint tint

    def style_run(run, font_name='Calibri', size_pt=11, color=COLOR_TEXT, bold=False, italic=False):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.color.rgb = color
        run.bold = bold
        run.italic = italic

    def set_cell_background(cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # Title & Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r = title_p.add_run('Charak Health Solutions · Dr. Surabhi Vaidya')
    style_run(r, 'Calibri', 22, COLOR_PRIMARY, bold=True)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    r = sub_p.add_run('Client User Acceptance Testing (UAT) Guide & Verification Checklist')
    style_run(r, 'Calibri', 13, COLOR_SECONDARY, italic=True)

    # Introduction Box
    intro_table = doc.add_table(rows=1, cols=1)
    intro_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = intro_table.rows[0].cells[0]
    set_cell_background(c, "F0FDF4")
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('Welcome to Your Testing Guide!\n')
    style_run(r, 'Calibri', 12, COLOR_PRIMARY, bold=True)
    r = p.add_run('This testing guide is prepared specifically for Dr. Surabhi Vaidya and clinic management. No technical knowledge is required. You will be testing every page address (slug), button, booking form, payment gateway, and staff portal to ensure the clinic website is completely ready for real patients.')
    style_run(r, 'Calibri', 10.5, COLOR_TEXT)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Master Directory Table
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run('1. Master Directory of Web Pages & Slugs')
    style_run(r, 'Calibri', 15, COLOR_PRIMARY, bold=True)

    slug_data = [
        ('/', 'Main Clinic Homepage', 'Comprehensive clinic landing page with non-surgical pain treatments, doctor credentials, and pricing.', 'All New Patients'),
        ('/#vsl', 'Video & Diagnostic Hero', 'Introduction video with Dr. Surabhi Vaidya and fast diagnostic test triggers.', 'First-time visitors'),
        ('/#flow', 'Treatment Pathways', 'Clear 3-step visual pathways: ₹201 Webinar, ₹1,000 Consultation, ₹4,000 Trial.', 'Patients seeking guidance'),
        ('/#symptoms', 'Symptom Recognition', 'Relatable pain struggles: Morning stiffness, stairs difficulty, shooting sciatica.', 'Hesitant patients'),
        ('/#science', '4-Pillar Ayurvedic Science', 'Clinical explanation of Marma therapy, Panchakarma herbal oils, joint lubrication.', 'Patients comparing with surgery'),
        ('/#doctor', 'Meet Dr. Surabhi Vaidya', 'Doctor qualifications (BAMS, MD Clinical Ayurveda), experience, and clinic vision.', 'Patients verifying credentials'),
        ('/#testimonials', 'Patient Success Stories', 'Real patient recovery video stories with one-click popup YouTube player.', 'Patients looking for proof'),
        ('/#pricing', 'Packages & Pricing', 'Transparent pricing for Consultation (₹1,000), 1-Day Trial (₹4,000), and Intensive Care.', 'Patients evaluating costs'),
        ('/#faq', 'Frequently Asked Questions', 'Expandable answers on relief speed, safety with diabetes, and clinic location.', 'Common patient doubts'),
        ('/?vsl=webinar', 'Webinar Video Version', 'Switches the hero video presentation specifically to Sunday Masterclass.', 'Sunday webinar ad traffic'),
        ('/webinar', 'Dedicated Webinar Page', 'Streamlined ad landing page dedicated to enrolling patients into the masterclass.', 'Ad campaigns (Meta/IG)'),
        ('/room', 'Live Masterclass Room', 'Live broadcast viewing room hosted on Webinar.gg with attendee intake.', 'Registered attendees'),
        ('/confirmation', 'Consultation Confirmation Page', 'Official post-booking pass showing clinic directions, what to bring, and tracking Google Ads conversions.', 'Patients who complete booking'),
        ('/about-us', 'About Us', 'Doctor credentials, clinical mission, 15+ years experience, and physical clinic details in Thane West.', 'Merchant & ads compliance'),
        ('/privacy-policy', 'Privacy Policy', 'Patient data protection, medical confidentiality, and PCI-DSS payment security disclosure.', 'Mandatory compliance (PayU/Google)'),
        ('/shipping-and-return-policy', 'Shipping & Delivery Policy', 'In-clinic appointment fulfillment, digital masterclass room delivery, and herbal dispatch standards.', 'Mandatory compliance for payment gateways'),
        ('/refund-policy', 'Refund & Cancellation Policy', '12-hour free reschedule/cancellation, 5-7 day bank refunds, and clear indication of business category.', 'Mandatory compliance for payment gateways'),
        ('/terms-and-conditions', 'Terms & Conditions', 'Clinical conditions of care, medical disclaimer, and governing jurisdiction.', 'Legal & merchant compliance'),
        ('/admin', 'Clinic Staff CRM Portal', 'Passcode-protected admin desk: View patient leads, send WhatsApp messages, edit settings.', 'Clinic staff & Dr. Surabhi')
    ]

    t = doc.add_table(rows=len(slug_data)+1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Slug / Address', 'Page Name', 'What It Is For', 'Primary Audience']
    for idx, text in enumerate(headers):
        cell = t.rows[0].cells[idx]
        set_cell_background(cell, "143026")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        style_run(r, 'Calibri', 10, RGBColor(255, 255, 255), bold=True)

    col_widths = [Inches(1.3), Inches(1.5), Inches(2.6), Inches(1.4)]
    for row_idx, data in enumerate(slug_data):
        row = t.rows[row_idx + 1]
        bg = COLOR_ZEBRA_BG if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            is_slug = (c_idx == 0)
            style_run(r, 'Consolas' if is_slug else 'Calibri', 9.5, COLOR_SECONDARY if is_slug else COLOR_TEXT, bold=is_slug)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 2. Pre-Testing Checklist
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run('2. Pre-Testing Checklist')
    style_run(r, 'Calibri', 15, COLOR_PRIMARY, bold=True)

    prep_items = [
        ("Test on a Smartphone", "Test on your iPhone or Android. Ensure buttons and text fit comfortably without sideways scrolling."),
        ("Test on a Laptop / Computer", "Verify widescreen layouts, readable typography, and the Staff CRM Portal."),
        ("Admin Login Passcode", "Use passcode: surabhi2026 (or click the green '1-Click Quick Unlock' button)."),
        ("Test Patient Information", "Use your personal name and mobile number to receive live WhatsApp booking confirmations.")
    ]

    for title, desc in prep_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f'✔ {title}: ')
        style_run(r, 'Calibri', 10.5, COLOR_PRIMARY, bold=True)
        r = p.add_run(desc)
        style_run(r, 'Calibri', 10.5, COLOR_TEXT)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. Step-by-Step Test Scenarios
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    r = h3.add_run('3. Step-by-Step Testing Checklist')
    style_run(r, 'Calibri', 15, COLOR_PRIMARY, bold=True)

    sections = [
        ("Section A: Main Clinic Homepage (/)", [
            ("Test 1.1: Top Clinic Notice Bar", "Check top green bar shows remaining slots and hours: Mon–Sat 9AM–8PM, Sun 10AM–2PM."),
            ("Test 1.2: Menu Navigation & Smooth Scrolling", "Click menu links (Overview, Pathways, Symptoms, Science, Doctor, Results, Packages, FAQ) to verify smooth glide."),
            ("Test 1.3: Video Player Popup", "Click any patient testimonial video card in '#testimonials' to verify YouTube video plays inside popup window."),
            ("Test 1.4: Expandable FAQ Accordion", "Click FAQ questions to verify answers expand and collapse smoothly."),
            ("Test 1.5: Clinic Address & Cal.id Link", "Check footer address (Cura 304, Raymond TenX Habitat, Pokharan Rd 2, Thane West) and Cal.id link."),
            ("Test 1.6: Mobile Sticky Bottom Bar", "On mobile phone, ensure 'Take 60-Second Pain Assessment' button stays docked comfortably at bottom.")
        ]),
        ("Section B: 60-Second Pain Diagnostic Assessment", [
            ("Test 2.1: Opening Assessment", "Click 'Start Assessment' button anywhere on page; modal opens with progress bar at 20%."),
            ("Test 2.2: Consultation Recommendation", "Select Knee Pain + Severe distress + Surgery advised + Fairly ready; quiz recommends In-Clinic Consultation (₹1,000)."),
            ("Test 2.3: Pay at Clinic Desk Option", "Fill patient name, mobile, preferred slot, click 'Pay at Clinic Desk'; confirmation screen displays with clinic directions."),
            ("Test 2.4: PayU Payment Gateway", "Click 'Pay ₹1,000 Online via PayU'; secure checkout window opens with UPI (GPay, PhonePe, Paytm), Cards, and NetBanking options."),
            ("Test 2.5: Cal.id Inline Booking Calendar", "Click 'Select Exact Time Slot on Calendar'; interactive calendar loads inside page without redirection."),
            ("Test 2.6: Live Masterclass Recommendation", "Select Mild pain + Hesitant/Skeptical; quiz smartly recommends Pathway 1: Live Hindi Masterclass (₹201).")
        ]),
        ("Section C: Dedicated Webinar Page (/webinar)", [
            ("Test 3.1: Dedicated Landing Cleanliness", "Open '/webinar' link directly; confirm clean design focused on webinar registration."),
            ("Test 3.2: Assessment & Ticket Booking", "Click registration button; assessment loads and directs to ₹201 ticket payment.")
        ]),
        ("Section D: Live Broadcast Room (/room)", [
            ("Test 4.1: Accessing the Live Room", "Open '/room'; verify dark-mode broadcast lobby with live status badge."),
            ("Test 4.2: Joining Stream with Mobile", "Enter patient name and mobile, click 'Join Live Broadcast Stream'; live player connects to Webinar.gg stream.")
        ]),
        ("Section E: Staff CRM Portal (/admin)", [
            ("Test 5.1: Passcode Login & 1-Click Unlock", "Visit '/admin'; unlock using passcode 'surabhi2026' or '1-Click Quick Unlock' button."),
            ("Test 5.2: Leads & CRM Intake (Tab 1)", "Check summary KPI counters and find the test patient record in the table."),
            ("Test 5.3: Real-Time Search & Status Filter", "Search by patient name; update lead status (e.g., to 'Consultation Scheduled')."),
            ("Test 5.4: Export Leads to CSV / Excel", "Click 'Export CSV'; confirm spreadsheet downloads with all patient intake details."),
            ("Test 5.5: WhatsApp Automation Status (Tab 2)", "Verify green badge 'Ready (Official BSP - AgentBandhu Cloud API)' and KPI delivery counters."),
            ("Test 5.6: WhatsApp Approved Templates", "Browse 10 Meta-approved templates; click 'Test Send' on any template to preview or send."),
            ("Test 5.7: Webinar Scheduler (Tab 3)", "View and update upcoming masterclass date/time and sync live attendee metrics."),
            ("Test 5.8: Clinic Configuration (Tab 4)", "Review clinic address, contact email, and Cal.id booking URL.")
        ])
    ]

    for sec_title, test_cases in sections:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(10)
        p_sec.paragraph_format.space_after = Pt(4)
        r = p_sec.add_run(sec_title)
        style_run(r, 'Calibri', 12, COLOR_PRIMARY, bold=True)

        for case_name, case_desc in test_cases:
            p_case = doc.add_paragraph()
            p_case.paragraph_format.space_before = Pt(2)
            p_case.paragraph_format.space_after = Pt(2)
            r = p_case.add_run(' [  ]  ')
            style_run(r, 'Consolas', 11, COLOR_MUTED, bold=True)
            r = p_case.add_run(f'{case_name}: ')
            style_run(r, 'Calibri', 10.5, COLOR_TEXT, bold=True)
            r = p_case.add_run(case_desc)
            style_run(r, 'Calibri', 10, COLOR_MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Feedback & Sign-Off Table
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)
    r = h4.add_run('4. Client Feedback & Sign-Off Sheet')
    style_run(r, 'Calibri', 15, COLOR_PRIMARY, bold=True)

    feedback_rows = [
        ('Clinic Hours & Address', '/', '[  ] PASS   [  ] REVIEW', ' '),
        ('Doctor Profile & Credentials', '/#doctor', '[  ] PASS   [  ] REVIEW', ' '),
        ('Patient Video Testimonials', '/#testimonials', '[  ] PASS   [  ] REVIEW', ' '),
        ('Packages & Pricing Table', '/#pricing', '[  ] PASS   [  ] REVIEW', ' '),
        ('60-Second Pain Assessment', 'Modal', '[  ] PASS   [  ] REVIEW', ' '),
        ('PayU Payment Gateway', 'Modal / Checkout', '[  ] PASS   [  ] REVIEW', ' '),
        ('Cal.id Inline Booking Calendar', 'Modal', '[  ] PASS   [  ] REVIEW', ' '),
        ('Live Broadcast Room', '/room', '[  ] PASS   [  ] REVIEW', ' '),
        ('Staff CRM Portal & Leads Table', '/admin (Tab 1)', '[  ] PASS   [  ] REVIEW', ' '),
        ('WhatsApp Messaging & Templates', '/admin (Tab 2)', '[  ] PASS   [  ] REVIEW', ' '),
        ('Webinar Scheduler & Metrics', '/admin (Tab 3)', '[  ] PASS   [  ] REVIEW', ' '),
        ('Clinic Settings & Cal.id Link', '/admin (Tab 4)', '[  ] PASS   [  ] REVIEW', ' ')
    ]

    fb_table = doc.add_table(rows=len(feedback_rows)+1, cols=4)
    fb_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fb_headers = ['Item Tested', 'Page / Slug', 'Status', 'Client Notes & Requested Updates']
    for idx, text in enumerate(fb_headers):
        cell = fb_table.rows[0].cells[idx]
        set_cell_background(cell, "143026")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        style_run(r, 'Calibri', 10, RGBColor(255, 255, 255), bold=True)

    col_widths = [Inches(1.8), Inches(1.3), Inches(1.5), Inches(2.2)]
    for row_idx, data in enumerate(feedback_rows):
        row = fb_table.rows[row_idx + 1]
        bg = COLOR_ZEBRA_BG if row_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            style_run(r, 'Calibri', 9.5, COLOR_TEXT)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Sign-off footer
    so_p = doc.add_paragraph()
    so_p.paragraph_format.space_before = Pt(6)
    so_p.paragraph_format.space_after = Pt(4)
    r = so_p.add_run('Client Review & Approval:\n')
    style_run(r, 'Calibri', 11, COLOR_PRIMARY, bold=True)
    r = so_p.add_run('Reviewed By: ___________________________    Date: ______________\nStatus:  [  ] Approved for Public Launch     [  ] Minor Updates Required')
    style_run(r, 'Calibri', 10.5, COLOR_TEXT)

    output_path = os.path.join('marketing', 'shareable-documents', 'UAT_Client_Testing_Guide.docx')
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f'Successfully generated {output_path}')

if __name__ == '__main__':
    create_uat_doc()
